from __future__ import annotations

import re
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table


WORKSPACE = Path(__file__).resolve().parents[1]
SOURCE_DOCX = WORKSPACE / "11023276-NguyenQuocViet-DoAn4-Tuan5.docx"
SOURCE_SQL = WORKSPACE / "doan4_daugia_tieng_viet.sql"
OUTPUT_DOCX = WORKSPACE / "11023276-NguyenQuocViet-DoAn4-Tuan5-hoan-thien-3.2.4.docx"


TABLE_PURPOSES = {
    "nguoi_dung": "Lưu tài khoản, vai trò và trạng thái hoạt động của người dùng.",
    "xac_minh_nguoi_ban": "Lưu hồ sơ và kết quả xác minh quyền bán.",
    "dia_chi_nguoi_dung": "Lưu các địa chỉ nhận hàng của người dùng.",
    "danh_muc": "Lưu danh mục, cấu hình thuộc tính động và chính sách kiểm định.",
    "san_pham": "Lưu hồ sơ sản phẩm do người bán tạo và trạng thái xét duyệt.",
    "kiem_dinh_san_pham": "Quản lý từng lần tiếp nhận, kiểm định và lưu giữ sản phẩm.",
    "phien_dau_gia": "Lưu cấu hình, thời gian, giá và trạng thái của phiên đấu giá.",
    "tham_gia_phien": "Lưu quan hệ theo dõi và mức giá tối đa bí mật của người tham gia.",
    "luot_tra_gia": "Lưu lịch sử các lượt giá công khai hợp lệ.",
    "dat_coc_dau_gia": "Quản lý toàn bộ vòng đời khoản cọc trước khi có đơn hàng.",
    "don_hang": "Lưu giao dịch sau đấu giá, tiền giữ trung gian và thông tin vận chuyển.",
    "thanh_toan": "Lưu các lần thanh toán phát sinh cho đơn hàng.",
    "de_nghi_mua_tiep_theo": "Lưu đề nghị Second Chance khi người thắng trước không thanh toán.",
    "tranh_chap": "Lưu hồ sơ, phản hồi và quyết định xử lý tranh chấp.",
    "tep_dinh_kem": "Lưu ảnh sản phẩm, bằng chứng và tài liệu kiểm định.",
    "danh_gia": "Lưu đánh giá hai chiều sau khi đơn hàng hoàn tất.",
    "yeu_cau_xu_ly": "Lưu yêu cầu hủy phiên và báo cáo sản phẩm.",
    "vi_pham": "Lưu hồ sơ vi phạm và quyết định xử lý của quản trị viên.",
    "thong_bao": "Lưu thông báo riêng theo từng người dùng và sự kiện.",
    "cau_hinh_he_thong": "Lưu các tham số nghiệp vụ có thể cấu hình.",
    "nhat_ky_hoat_dong": "Lưu vết thao tác, thay đổi dữ liệu và khóa chống lặp.",
}


TABLE_CAPTION_NAMES = {
    "nguoi_dung": "người dùng",
    "xac_minh_nguoi_ban": "xác minh người bán",
    "dia_chi_nguoi_dung": "địa chỉ người dùng",
    "danh_muc": "danh mục",
    "san_pham": "sản phẩm",
    "kiem_dinh_san_pham": "kiểm định sản phẩm",
    "phien_dau_gia": "phiên đấu giá",
    "tham_gia_phien": "tham gia phiên",
    "luot_tra_gia": "lượt trả giá",
    "dat_coc_dau_gia": "đặt cọc đấu giá",
    "don_hang": "đơn hàng",
    "thanh_toan": "thanh toán",
    "de_nghi_mua_tiep_theo": "đề nghị mua tiếp theo",
    "tranh_chap": "tranh chấp",
    "tep_dinh_kem": "tệp đính kèm",
    "danh_gia": "đánh giá",
    "yeu_cau_xu_ly": "yêu cầu xử lý",
    "vi_pham": "vi phạm",
    "thong_bao": "thông báo",
    "cau_hinh_he_thong": "cấu hình hệ thống",
    "nhat_ky_hoat_dong": "nhật ký hoạt động",
}


GROUPS = [
    ("a) Nhóm tài khoản", ["nguoi_dung", "xac_minh_nguoi_ban", "dia_chi_nguoi_dung"]),
    ("b) Nhóm danh mục, sản phẩm và kiểm định", ["danh_muc", "san_pham", "kiem_dinh_san_pham"]),
    ("c) Nhóm đấu giá và đặt cọc", ["phien_dau_gia", "tham_gia_phien", "luot_tra_gia", "dat_coc_dau_gia"]),
    ("d) Nhóm đơn hàng và thanh toán", ["don_hang", "thanh_toan", "de_nghi_mua_tiep_theo"]),
    ("e) Nhóm hậu giao dịch", ["tranh_chap", "tep_dinh_kem", "danh_gia"]),
    ("f) Nhóm hỗ trợ và quản trị", ["yeu_cau_xu_ly", "vi_pham", "thong_bao", "cau_hinh_he_thong", "nhat_ky_hoat_dong"]),
]


SPECIAL_DESCRIPTIONS = {
    "id": "Khóa chính định danh duy nhất của bản ghi.",
    "mat_khau_bam": "Mật khẩu đã được băm trước khi lưu.",
    "anh_dai_dien": "Đường dẫn ảnh đại diện của người dùng.",
    "vai_tro": "Vai trò truy cập của tài khoản.",
    "trang_thai_nguoi_ban": "Trạng thái đăng ký và xác minh quyền bán.",
    "trang_thai_tai_khoan": "Trạng thái hoạt động, tạm ngưng hoặc khóa tài khoản.",
    "cau_hinh_thuoc_tinh": "Mảng JSON mô tả các thuộc tính động của danh mục.",
    "thuoc_tinh_json": "Đối tượng JSON lưu giá trị thuộc tính của sản phẩm.",
    "bat_buoc_kiem_dinh": "Giá trị chính sách kiểm định được chụp khi sản phẩm gửi duyệt.",
    "ngay_chup_chinh_sach_kiem_dinh": "Thời điểm chụp chính sách kiểm định cho sản phẩm.",
    "gia_toi_da": "Mức giá tối đa bí mật người tham gia cam kết trả.",
    "thoi_gian_dat_gia_toi_da": "Thời điểm ghi nhận mức giá tối đa, dùng để xác định ưu tiên.",
    "gia_hien_tai": "Mức giá công khai hiện tại của phiên.",
    "gia_san_pham": "Giá sản phẩm được chốt khi tạo đơn.",
    "phi_van_chuyen": "Phí vận chuyển được áp dụng cho giao dịch.",
    "tong_tien": "Tổng tiền của đơn, gồm giá sản phẩm và phí vận chuyển.",
    "gia_san": "Mức giá tối thiểu để phiên có người thắng.",
    "gia_mua_ngay": "Mức giá cố định dùng cho chức năng Mua ngay.",
    "dat_gia_san": "Cờ cho biết giá công khai đã đạt giá sàn.",
    "thoi_gian_ket_thuc_goc": "Thời điểm kết thúc trước khi áp dụng gia hạn.",
    "thoi_gian_ket_thuc": "Thời điểm kết thúc hiện tại sau các lần gia hạn.",
    "nguong_phut_chot_giay": "Số giây cuối phiên có thể kích hoạt gia hạn.",
    "so_giay_gia_han": "Số giây cộng thêm khi gia hạn phiên.",
    "yeu_cau_dat_coc": "Cờ xác định phiên có bắt buộc đặt cọc hay không.",
    "so_tien_dat_coc": "Số tiền cọc được chụp tại thời điểm tạo phiên.",
    "khoa_yeu_cau": "Khóa idempotency dùng để chống xử lý lặp yêu cầu tiền.",
    "so_tien": "Số tiền của giao dịch hoặc khoản nghiệp vụ tương ứng.",
    "trang_thai": "Trạng thái hiện tại của bản ghi.",
    "ly_do": "Lý do được ghi nhận cho yêu cầu hoặc tình huống nghiệp vụ.",
    "nguon_don": "Nguồn hình thành đơn: thắng đấu giá hoặc Second Chance.",
    "trang_thai_giu_tien": "Trạng thái đối soát khoản tiền đang được giữ trung gian.",
    "so_tien_dang_giu": "Số tiền sinh tự động bằng số đã thu trừ số đã hoàn và đã giải ngân.",
    "du_lieu_lich_su": "Cờ nhận biết bản ghi được chuyển đổi từ cấu trúc dữ liệu cũ.",
    "giu_tien_id_cu": "Mã bản ghi giữ tiền cũ, được bảo toàn để đối chiếu migration.",
    "giu_tien_ngay_tao": "Thời điểm tạo của bản ghi giữ tiền cũ.",
    "giu_tien_ngay_cap_nhat": "Thời điểm cập nhật của bản ghi giữ tiền cũ.",
    "van_chuyen_id_cu": "Mã bản ghi vận chuyển cũ, được bảo toàn để đối chiếu migration.",
    "van_chuyen_ngay_tao": "Thời điểm tạo của bản ghi vận chuyển cũ.",
    "van_chuyen_ngay_cap_nhat": "Thời điểm cập nhật của bản ghi vận chuyển cũ.",
    "phien_con_nghia_vu": "Cột sinh dùng để bảo đảm mỗi phiên chỉ có một đơn còn nghĩa vụ.",
    "so_tien_con_phai_thanh_toan": "Số tiền sinh tự động bằng tổng tiền trừ số tiền đã thu.",
    "nguon_gui_hang": "Nguồn thực hiện gửi hàng: người bán hoặc trung tâm.",
    "don_da_thu_tien": "Cột sinh dùng để ngăn tạo nhiều kết quả thu tiền cho cùng đơn.",
    "phien_dang_de_nghi": "Cột sinh dùng để bảo đảm mỗi phiên chỉ có một đề nghị đang chờ.",
    "don_dang_tranh_chap": "Cột sinh dùng để ngăn một đơn có nhiều tranh chấp đang mở.",
    "san_pham_anh_chinh": "Cột sinh dùng để bảo đảm mỗi sản phẩm chỉ có một ảnh chính.",
    "phien_huy_dang_mo": "Cột sinh dùng để ngăn yêu cầu hủy phiên đang mở bị tạo trùng.",
    "san_pham_bao_cao_dang_mo": "Cột sinh dùng để ngăn báo cáo sản phẩm đang mở bị tạo trùng.",
    "khoa_su_kien": "Khóa sự kiện dùng để tránh tạo thông báo trùng.",
    "ma_yeu_cau": "Mã yêu cầu duy nhất dùng cho truy vết và chống xử lý lặp.",
    "du_lieu_cu": "Ảnh chụp JSON của dữ liệu trước khi thay đổi.",
    "du_lieu_moi": "Ảnh chụp JSON của dữ liệu sau khi thay đổi.",
}


PHRASE_LABELS = {
    "nguoi_dung": "người dùng",
    "nguoi_ban": "người bán",
    "nguoi_mua": "người mua",
    "nguoi_nhan": "người nhận",
    "nguoi_duyet": "người duyệt",
    "nguoi_xu_ly": "người xử lý",
    "nguoi_yeu_cau": "người yêu cầu",
    "nguoi_thuc_hien": "người thực hiện",
    "nguoi_cap_nhat": "người cập nhật",
    "nguoi_tra_gia": "người trả giá",
    "nguoi_danh_gia": "người đánh giá",
    "nguoi_duoc_danh_gia": "người được đánh giá",
    "nguoi_dan_dau": "người dẫn đầu",
    "phien_dau_gia": "phiên đấu giá",
    "luot_tra_gia": "lượt trả giá",
    "don_hang": "đơn hàng",
    "san_pham": "sản phẩm",
    "danh_muc": "danh mục",
    "kiem_dinh": "kiểm định",
    "tranh_chap": "tranh chấp",
    "thanh_toan": "thanh toán",
    "van_chuyen": "vận chuyển",
    "giai_ngan": "giải ngân",
    "chuyen_gia": "chuyên gia",
    "tinh_trang": "tình trạng",
    "tinh_thanh": "tỉnh thành",
    "dia_chi": "địa chỉ",
    "duong_dan": "đường dẫn",
    "trang_thai": "trạng thái",
    "thoi_gian": "thời gian",
    "so_tien": "số tiền",
    "gia_tri": "giá trị",
    "gia_khoi_diem": "giá khởi điểm",
    "gia_mua_ngay": "giá Mua ngay",
    "gia_de_nghi": "giá đề nghị",
    "gia_san_pham": "giá sản phẩm",
    "gia_toi_da": "giá tối đa",
    "gia_hien_tai": "giá hiện tại",
    "gia_san": "giá sàn",
    "ngay_cap_nhat": "ngày cập nhật",
    "ngay_tao": "ngày tạo",
    "giay_to": "giấy tờ",
    "don_vi": "đơn vị",
    "ly_do": "lý do",
    "ghi_chu": "ghi chú",
    "yeu_cau": "yêu cầu",
}


TOKEN_LABELS = {
    "id": "ID", "ho": "họ", "ten": "tên", "email": "email", "mat": "mật", "khau": "khẩu",
    "bam": "băm", "so": "số", "dien": "điện", "thoai": "thoại", "anh": "ảnh", "dai": "đại",
    "vai": "vai", "tro": "trò", "trang": "trạng", "thai": "thái", "tai": "tài", "khoan": "khoản",
    "ngay": "ngày", "xac": "xác", "minh": "minh", "lan": "lần", "dang": "đang", "nhap": "nhập",
    "cuoi": "cuối", "loai": "loại", "giay": "giấy", "to": "tờ", "truoc": "trước", "sau": "sau",
    "ngan": "ngân", "hang": "hàng", "chu": "chủ", "gui": "gửi", "duyet": "duyệt", "tinh": "tỉnh",
    "thanh": "thành", "quan": "quận", "huyen": "huyện", "phuong": "phường", "xa": "xã", "chi": "chi",
    "tiet": "tiết", "la": "là", "mac": "mặc", "dinh": "định", "cha": "cha", "duong": "đường",
    "dan": "dẫn", "mo": "mô", "ta": "tả", "thu": "thứ", "tu": "tự", "cau": "cấu", "hinh": "hình",
    "thuoc": "thuộc", "bat": "bắt", "buoc": "buộc", "tieu": "tiêu", "de": "đề", "thuong": "thương",
    "hieu": "hiệu", "ly": "lý", "do": "do", "chinh": "chính", "sach": "sách", "ma": "mã",
    "trung": "trung", "tam": "tâm", "van": "vận", "don": "đơn", "den": "đến", "nhan": "nhận",
    "kien": "kiện", "ghi": "ghi", "chuyen": "chuyển", "gia": "giá", "ket": "kết", "qua": "quả",
    "chung": "chứng", "tra": "trả", "roi": "rời", "cap": "cập", "khoi": "khởi", "diem": "điểm",
    "san": "sàn", "mua": "mua", "cho": "cho", "phep": "phép", "hien": "hiện", "dau": "đấu",
    "thoi": "thời", "gian": "gian", "thuc": "thức", "goc": "gốc", "dat": "đặt", "chong": "chống",
    "phut": "phút", "chot": "chót", "nguong": "ngưỡng", "giay": "giây", "han": "hạn", "tong": "tổng",
    "luot": "lượt", "phi": "phí", "coc": "cọc", "tham": "tham", "theo": "theo", "doi": "dõi",
    "toi": "tối", "da": "đã", "tien": "tiền", "truc": "trực", "phuong": "phương", "giao": "giao",
    "dich": "dịch", "khoa": "khóa", "hoan": "hoàn", "vao": "vào", "khong": "không", "xu": "xử",
    "nguon": "nguồn", "thang": "thắng", "kiem": "kiểm", "tat": "tất", "huy": "hủy", "sdt": "SĐT",
    "giu": "giữ", "du": "dữ", "lieu": "liệu", "lich": "lịch", "su": "sử", "moc": "mốc",
    "khieu": "khiếu", "nai": "nại", "chua": "chưa", "can": "cần", "admin": "Admin", "cu": "cũ",
    "con": "còn", "nghia": "nghĩa", "vu": "vụ", "the": "thẻ", "het": "hết", "nghi": "nghị",
    "tiep": "tiếp", "phan": "phản", "hoi": "hồi", "chap": "chấp", "tep": "tệp", "bang": "bằng",
    "noi": "nội", "dung": "dung", "sao": "sao", "xet": "xét", "bao": "báo", "cao": "cáo",
    "pham": "phạm", "thong": "thông", "lien": "liên", "doc": "đọc", "kieu": "kiểu", "hanh": "hành",
    "dong": "động", "doi": "đối", "tuong": "tượng", "ip": "IP", "trinh": "trình", "thiet": "thiết",
    "bi": "bị", "selfie": "selfie", "serial": "serial", "web": "Web",
}


def split_sql_items(body: str) -> list[str]:
    items: list[str] = []
    start = 0
    depth = 0
    quote: str | None = None
    escaped = False
    for index, char in enumerate(body):
        if quote:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == quote:
                quote = None
        else:
            if char in "'\"":
                quote = char
            elif char == "(":
                depth += 1
            elif char == ")":
                depth -= 1
            elif char == "," and depth == 0:
                items.append(body[start:index].strip())
                start = index + 1
    items.append(body[start:].strip())
    return items


def parse_schema(sql: str):
    table_pattern = re.compile(r"CREATE TABLE `([^`]+)` \((.*?)\)\s*ENGINE=", re.S)
    tables = {}
    unique_columns: dict[str, dict[str, str]] = defaultdict(dict)
    primary_columns: dict[str, set[str]] = defaultdict(set)

    for table_name, body in table_pattern.findall(sql):
        columns = []
        for item in split_sql_items(body):
            match = re.match(r"`([^`]+)`\s+(.+)", item, re.S)
            if match:
                columns.append({"name": match.group(1), "definition": re.sub(r"\s+", " ", match.group(2)).strip()})
                continue

            primary_match = re.match(r"PRIMARY KEY\s*\(([^)]+)\)", item, re.I)
            if primary_match:
                primary_columns[table_name].update(re.findall(r"`([^`]+)`", primary_match.group(1)))
                continue

            unique_match = re.match(r"UNIQUE KEY\s+`[^`]+`\s*\(([^)]+)\)", item, re.I)
            if unique_match:
                names = re.findall(r"`([^`]+)`", unique_match.group(1))
                kind = "UNIQUE" if len(names) == 1 else "UNIQUE ghép"
                for name in names:
                    unique_columns[table_name][name] = kind

        tables[table_name] = columns

    foreign_keys: dict[tuple[str, str], tuple[str, str]] = {}
    alter_pattern = re.compile(r"ALTER TABLE `([^`]+)`(.*?);", re.S)
    fk_pattern = re.compile(
        r"FOREIGN KEY \(`([^`]+)`\) REFERENCES `([^`]+)` \(`([^`]+)`\)",
        re.S,
    )
    for table_name, alter_body in alter_pattern.findall(sql):
        for column_name, target_table, target_column in fk_pattern.findall(alter_body):
            foreign_keys[(table_name, column_name)] = (target_table, target_column)

    return tables, primary_columns, unique_columns, foreign_keys


def display_type(definition: str) -> tuple[str, list[str]]:
    enum_match = re.match(r"enum\((.*?)\)", definition, re.I | re.S)
    if enum_match:
        return "ENUM", re.findall(r"'([^']+)'", enum_match.group(1))

    patterns = [
        r"decimal\(\d+\s*,\s*\d+\)", r"varchar\(\d+\)", r"tinyint\(\d+\)(?: unsigned)?",
        r"bigint(?: unsigned)?", r"tinyint(?: unsigned)?", r"int(?: unsigned)?", r"timestamp", r"datetime",
        r"text", r"json",
    ]
    for pattern in patterns:
        match = re.match(pattern, definition, re.I)
        if match:
            return re.sub(r"\s+", "", match.group(0)).upper().replace("UNSIGNED", " UNSIGNED"), []
    return definition.split()[0].upper(), []


def default_value(definition: str) -> str | None:
    match = re.search(r"\bDEFAULT\s+(CURRENT_TIMESTAMP|'[^']*'|NULL|[^\s,]+)", definition, re.I)
    if not match or match.group(1).upper() == "NULL":
        return None
    return match.group(1).strip("'")


def constraints_for(table_name, column_name, definition, primary_columns, unique_columns, foreign_keys) -> str:
    parts = []
    if column_name in primary_columns[table_name]:
        parts.append("PK")
    if (table_name, column_name) in foreign_keys:
        target_table, target_column = foreign_keys[(table_name, column_name)]
        parts.append(f"FK → {target_table}({target_column})")
    if column_name in unique_columns[table_name]:
        parts.append(unique_columns[table_name][column_name])
    if "NOT NULL" in definition.upper():
        parts.append("NOT NULL")
    else:
        parts.append("Cho phép NULL")
    if "AUTO_INCREMENT" in definition.upper():
        parts.append("Tự tăng")
    if "GENERATED ALWAYS" in definition.upper():
        parts.append("Cột sinh")
    value = default_value(definition)
    if value is not None:
        parts.append(f"Mặc định {value}")
    if "ON UPDATE CURRENT_TIMESTAMP" in definition.upper():
        parts.append("Tự cập nhật")
    return ", ".join(parts)


def human_label(column_name: str) -> str:
    text = column_name
    placeholders = {}
    for index, phrase in enumerate(sorted(PHRASE_LABELS, key=len, reverse=True)):
        marker = f"zz{index}zz"
        if phrase in text:
            text = text.replace(phrase, marker)
            placeholders[marker] = PHRASE_LABELS[phrase]

    words = []
    for token in text.split("_"):
        words.append(placeholders.get(token, TOKEN_LABELS.get(token, token)))
    label = " ".join(words)
    for marker, value in placeholders.items():
        label = label.replace(marker, value)
    return label.strip()


def description_for(table_name: str, column_name: str, definition: str, foreign_keys) -> str:
    if column_name in SPECIAL_DESCRIPTIONS:
        description = SPECIAL_DESCRIPTIONS[column_name]
    elif (table_name, column_name) in foreign_keys:
        target_table, _ = foreign_keys[(table_name, column_name)]
        description = f"Tham chiếu đến bản ghi tương ứng trong bảng {target_table}."
    else:
        label = human_label(column_name)
        if column_name.startswith("ngay_") or column_name.startswith("thoi_gian_"):
            description = f"Thời điểm {label.removeprefix('ngày ').removeprefix('thời gian ')}."
        elif column_name.startswith("han_") or column_name.endswith("_luc"):
            description = f"Thời hạn {label.removeprefix('hạn ')}."
        elif column_name.startswith("ma_"):
            description = f"Mã {label.removeprefix('mã ')} phục vụ nhận dạng và tra cứu."
        elif column_name.startswith("trang_thai"):
            detail = label.removeprefix("trạng thái ")
            description = f"Trạng thái {detail} của bản ghi." if detail else "Trạng thái hiện tại của bản ghi."
        elif column_name.startswith("ly_do"):
            detail = label.removeprefix("lý do ")
            description = f"Lý do {detail}." if detail else "Lý do liên quan đến bản ghi."
        elif column_name.startswith("so_tien"):
            detail = label.removeprefix("số tiền ")
            description = f"Số tiền {detail}." if detail else "Số tiền của bản ghi."
        elif column_name.startswith("gia_") or column_name.startswith("phi_"):
            description = label[:1].upper() + label[1:] + "."
        elif column_name.startswith("la_") or column_name.startswith("co_") or column_name.startswith("can_") or column_name.startswith("dang_") or column_name.startswith("cho_phep_") or column_name.startswith("bat_") or column_name.startswith("yeu_cau_"):
            description = f"Cờ xác định {label}."
        else:
            description = f"Lưu {label}."

    _, enum_values = display_type(definition)
    if enum_values:
        description += " Giá trị: " + ", ".join(enum_values) + "."
    return description


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="000000", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def format_cell(cell, text, *, header=False, center=False, shade=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = "Nội dung"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(text))
    if header:
        run.style = "Strong"
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_borders(cell)
    if shade:
        set_cell_shading(cell, shade)


def configure_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])


def add_summary_table(document, target, tables):
    caption = document.add_paragraph(style="Bang")
    caption.paragraph_format.keep_with_next = True
    caption.add_run("Bảng 3-33: Danh sách các bảng trong cơ sở dữ liệu")
    target._p.addprevious(caption._p)

    summary = document.add_table(rows=1, cols=3)
    widths = [0.45, 1.9, 3.88]
    headers = ["STT", "Tên bảng", "Chức năng"]
    for index, header in enumerate(headers):
        format_cell(summary.rows[0].cells[index], header, header=True, center=True, shade="A5C9EB")
    set_repeat_table_header(summary.rows[0])

    for number, table_name in enumerate(tables, 1):
        cells = summary.add_row().cells
        values = [number, table_name, TABLE_PURPOSES[table_name]]
        for index, value in enumerate(values):
            format_cell(cells[index], value, center=index == 0)
    configure_table(summary, widths)
    target._p.addprevious(summary._tbl)


def add_detail_table(document, target, table_number, table_name, columns, primary_columns, unique_columns, foreign_keys):
    caption = document.add_paragraph(style="Bang")
    caption.paragraph_format.keep_with_next = True
    caption.add_run(f"Bảng 3-{table_number}: Cấu trúc bảng {table_name}")
    target._p.addprevious(caption._p)

    table = document.add_table(rows=1, cols=5)
    widths = [0.38, 1.23, 1.03, 1.65, 1.94]
    headers = ["STT", "Tên trường", "Kiểu dữ liệu", "Ràng buộc", "Mô tả"]
    for index, header in enumerate(headers):
        format_cell(table.rows[0].cells[index], header, header=True, center=True, shade="A5C9EB")
    set_repeat_table_header(table.rows[0])

    for number, column in enumerate(columns, 1):
        type_name, _ = display_type(column["definition"])
        constraints = constraints_for(
            table_name,
            column["name"],
            column["definition"],
            primary_columns,
            unique_columns,
            foreign_keys,
        )
        description = description_for(table_name, column["name"], column["definition"], foreign_keys)
        values = [number, column["name"], type_name, constraints, description]
        row = table.add_row()
        for index, value in enumerate(values):
            format_cell(
                row.cells[index],
                value,
                center=index in (0, 2),
            )
    configure_table(table, widths)
    target._p.addprevious(table._tbl)


def add_update_fields_setting(document):
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def clear_paragraph_content(paragraph):
    paragraph.clear()
    return paragraph


def build_table_from_report_template(
    document,
    target,
    template_table,
    table_name,
    columns,
    primary_columns,
    unique_columns,
    foreign_keys,
):
    template_xml = template_table._tbl
    if len(template_xml.tr_lst) < 2:
        raise RuntimeError("Bảng mẫu không có hàng dữ liệu để sao chép định dạng")

    new_table_xml = deepcopy(template_xml)
    body_row_template = deepcopy(template_xml.tr_lst[1])
    for row_xml in list(new_table_xml.tr_lst[1:]):
        new_table_xml.remove(row_xml)

    wrapper = Table(new_table_xml, document._body)
    set_repeat_table_header(wrapper.rows[0])

    for number, column in enumerate(columns, 1):
        new_table_xml.append(deepcopy(body_row_template))
        wrapper = Table(new_table_xml, document._body)
        row = wrapper.rows[-1]
        type_name, _ = display_type(column["definition"])
        values = [
            str(number),
            column["name"],
            type_name,
            constraints_for(
                table_name,
                column["name"],
                column["definition"],
                primary_columns,
                unique_columns,
                foreign_keys,
            ),
            description_for(table_name, column["name"], column["definition"], foreign_keys),
        ]
        for cell, value in zip(row.cells, values):
            paragraph = clear_paragraph_content(cell.paragraphs[0])
            paragraph.add_run(value)

    target._p.addprevious(new_table_xml)


def main():
    sql = SOURCE_SQL.read_text(encoding="utf-8")
    tables, primary_columns, unique_columns, foreign_keys = parse_schema(sql)

    expected_tables = [name for _, names in GROUPS for name in names]
    if list(tables) != expected_tables:
        raise RuntimeError(f"Danh sách bảng SQL không khớp: {list(tables)}")

    document = Document(str(SOURCE_DOCX))
    heading = next((p for p in document.paragraphs if p.text.strip() == "Thiết kế cơ sở dữ liệu"), None)
    figure_caption = next((p for p in document.paragraphs if p.text.strip() == "Hình Sơ đồ quan hệ cơ sở dữ liệu"), None)
    target = next((p for p in document.paragraphs if p.text.strip() == "Thiết kế giao diện"), None)
    if heading is None or figure_caption is None or target is None:
        raise RuntimeError("Không tìm thấy đầy đủ mục Thiết kế cơ sở dữ liệu, chú thích sơ đồ hoặc mục Thiết kế giao diện")

    sample_table = next(
        (
            table
            for table in document.tables
            if len(table.columns) == 5
            and [cell.text.strip() for cell in table.rows[0].cells]
            == ["STT", "Tên trường", "Kiểu dữ liệu", "Ràng buộc", "Ý nghĩa"]
        ),
        None,
    )
    if sample_table is None:
        raise RuntimeError("Không tìm thấy bảng mẫu thiết kế cơ sở dữ liệu trong báo cáo")

    heading.text = "3.2.4 Thiết kế cơ sở dữ liệu"

    next_node = figure_caption._p.getnext()
    while next_node is not None and next_node is not target._p:
        following = next_node.getnext()
        next_node.getparent().remove(next_node)
        next_node = following

    table_number = 33
    for table_name in expected_tables:
        label = document.add_paragraph(style="Nội dung")
        label.paragraph_format.keep_with_next = True
        label.add_run(f"Bảng {table_name}")
        target._p.addprevious(label._p)

        caption = document.add_paragraph(style="Bang")
        caption.paragraph_format.keep_with_next = True
        caption.add_run(f"Bảng 3-{table_number}: Cấu trúc bảng {TABLE_CAPTION_NAMES[table_name]}")
        target._p.addprevious(caption._p)

        build_table_from_report_template(
            document,
            target,
            sample_table,
            table_name,
            tables[table_name],
            primary_columns,
            unique_columns,
            foreign_keys,
        )
        table_number += 1

    add_update_fields_setting(document)
    document.core_properties.title = "Xây dựng hệ thống đấu giá trực tuyến"
    document.save(str(OUTPUT_DOCX))
    print(f"Đã tạo: {OUTPUT_DOCX}")
    print(f"Số bảng: {len(tables)}; số trường: {sum(len(columns) for columns in tables.values())}")


if __name__ == "__main__":
    main()
